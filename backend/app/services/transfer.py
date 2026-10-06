"""CSV import (with dry run) and CSV / XLSX export."""

import csv
import io
import re
from collections.abc import Sequence
from datetime import date
from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.clock import today
from app.deadline import DeadlineStatus
from app.errors import AppError
from app.fmt import format_date, parse_date
from app.models import Category, Person, Project, Stage
from app.schemas import ImportFieldError, ImportReport, ImportRow, ProjectCreate
from app.services.projects import insert_project

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
DEFAULT_STAGE = "Not Started"

STATUS_LABELS = {
    DeadlineStatus.done: "Done",
    DeadlineStatus.no_deadline: "No Deadline",
    DeadlineStatus.overdue: "Overdue",
    DeadlineStatus.due_soon: "Due Soon",
    DeadlineStatus.on_track: "On Track",
}

# header (normalised) -> field name
HEADER_ALIASES = {
    "category": "category",
    "project": "name",
    "name": "name",
    "description": "description",
    "ongoing issue": "ongoing_issue",
    "stage": "stage",
    "status notes": "status_notes",
    "plan date": "plan_date",
    "adjusted plan date": "adjusted_plan_date",
    "actual date": "actual_date",
    "projected start": "projected_start_date",
    "projected start date": "projected_start_date",
    "deadline note": "deadline_note",
    "pic": "pics",
    "pics": "pics",
}
EXPORT_COLUMNS = [
    "Category",
    "Project",
    "Description",
    "Ongoing Issue",
    "Stage",
    "Status Notes",
    "Plan Date",
    "Adjusted Plan Date",
    "Actual Date",
    "Projected Start",
    "Deadline Note",
    "PIC",
    "Effective Deadline",
    "Deadline Status",
]
# Export-only columns that are accepted (and ignored) on import so exports round-trip.
DERIVED_HEADERS = {"effective deadline", "deadline status"}
DATE_FIELDS = ("plan_date", "adjusted_plan_date", "actual_date", "projected_start_date")


def _normalise_header(header: str) -> str:
    return re.sub(r"[\s_]+", " ", header.strip().lstrip("﻿").lower())


def _clean_cell(value: str | None) -> str:
    text = (value or "").strip()
    if len(text) > 1 and text[0] == "'" and text[1] in "=+-@":
        return text[1:]  # undo the formula-injection guard added on export
    return text


def _safe_cell(value: str) -> str:
    """Neutralise spreadsheet formulas in exported text."""
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def _row_values(p: Project, today_: date) -> list[str | date | None]:
    from app.deadline import deadline_status, effective_deadline

    status = deadline_status(p.actual_date, p.plan_date, p.adjusted_plan_date, today_)
    return [
        p.category.value,
        p.name,
        p.description,
        p.ongoing_issue,
        p.stage.name,
        p.status_notes,
        p.plan_date,
        p.adjusted_plan_date,
        p.actual_date,
        p.projected_start_date,
        p.deadline_note,
        ", ".join(x.name for x in p.pics) or "TBD",
        effective_deadline(p.plan_date, p.adjusted_plan_date),
        STATUS_LABELS[status],
    ]


def export_csv(projects: Sequence[Project]) -> bytes:
    today_ = today()
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_COLUMNS)
    for p in projects:
        writer.writerow(
            [
                format_date(v) if isinstance(v, date) or v is None else _safe_cell(v)
                for v in _row_values(p, today_)
            ]
        )
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def export_xlsx(projects: Sequence[Project]) -> bytes:
    today_ = today()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Projects"
    ws.append(EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"
    for p in projects:
        ws.append([_safe_cell(v) if isinstance(v, str) else v for v in _row_values(p, today_)])
    date_columns = {
        EXPORT_COLUMNS.index(c) + 1
        for c in EXPORT_COLUMNS
        if "Date" in c or c == "Projected Start" or c == "Effective Deadline"
    }
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in date_columns:
                cell.number_format = "ddd, mmm dd, yy"
                if cell.value is None:
                    cell.value = "TBD"
    widths = [12, 36, 50, 40, 18, 50, 18, 20, 18, 18, 24, 22, 20, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _decode(raw: bytes) -> str:
    if len(raw) > MAX_IMPORT_BYTES:
        raise AppError(413, "file_too_large", "CSV file is larger than 2 MB")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AppError(400, "invalid_csv", "File is not valid UTF-8 text") from exc


def _read_csv(raw: bytes) -> tuple[list[dict[str, str]], list[int], list[str]]:
    text = _decode(raw)
    if not text.strip():
        raise AppError(400, "invalid_csv", "The file is empty")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = next(reader)
    except StopIteration as exc:
        raise AppError(400, "invalid_csv", "The file is empty") from exc
    fields: list[str | None] = []
    ignored: list[str] = []
    for h in header:
        norm = _normalise_header(h)
        if norm in HEADER_ALIASES:
            fields.append(HEADER_ALIASES[norm])
        else:
            fields.append(None)
            if norm and norm not in DERIVED_HEADERS:
                ignored.append(h.strip())
    named = [f for f in fields if f]
    if len(named) != len(set(named)):
        raise AppError(400, "invalid_csv", "The header row contains duplicate columns")
    missing = [c for c in ("name", "category") if c not in named]
    if missing:
        label = {"name": "Project", "category": "Category"}
        raise AppError(
            400,
            "invalid_csv",
            "Missing required column(s): " + ", ".join(label[m] for m in missing),
        )
    rows: list[dict[str, str]] = []
    row_numbers: list[int] = []
    try:
        for line_no, cells in enumerate(reader, start=2):
            if not any(c.strip() for c in cells):
                continue
            rows.append(
                {
                    f: _clean_cell(cells[i] if i < len(cells) else "")
                    for i, f in enumerate(fields)
                    if f
                }
            )
            row_numbers.append(line_no)
    except csv.Error as exc:
        raise AppError(400, "invalid_csv", f"Could not parse CSV: {exc}") from exc
    if len(rows) > MAX_IMPORT_ROWS:
        raise AppError(400, "too_many_rows", f"At most {MAX_IMPORT_ROWS} rows can be imported")
    return rows, row_numbers, ignored


def _validate(
    session: Session, raw: bytes
) -> tuple[list[ImportRow], list[tuple[ProjectCreate, Stage, list[Person]] | None], list[str]]:
    rows, row_numbers, ignored = _read_csv(raw)

    stages = {s.name.lower(): s for s in session.scalars(sa.select(Stage))}
    people = {p.name.lower(): p for p in session.scalars(sa.select(Person))}
    existing = {
        n.lower()
        for n in session.scalars(sa.select(Project.name).where(Project.deleted_at.is_(None)))
    }
    seen: dict[str, int] = {}

    report_rows: list[ImportRow] = []
    parsed: list[tuple[ProjectCreate, Stage, list[Person]] | None] = []
    for row_no, raw_row in zip(row_numbers, rows, strict=True):
        errors: list[ImportFieldError] = []
        payload: dict[str, Any] = {
            "name": raw_row.get("name", ""),
            "description": raw_row.get("description", ""),
            "status_notes": raw_row.get("status_notes", ""),
            "ongoing_issue": raw_row.get("ongoing_issue", ""),
            "deadline_note": raw_row.get("deadline_note", ""),
        }

        category = raw_row.get("category", "").lower()
        if category in {c.value for c in Category}:
            payload["category"] = category
        else:
            errors.append(
                ImportFieldError(
                    field="category",
                    message="Category must be current, pending or maintenance",
                )
            )

        stage_name = raw_row.get("stage", "") or DEFAULT_STAGE
        stage = stages.get(stage_name.lower())
        if stage is None:
            errors.append(ImportFieldError(field="stage", message=f"Unknown stage '{stage_name}'"))
        else:
            payload["stage_id"] = stage.id

        for name in DATE_FIELDS:
            try:
                payload[name] = parse_date(raw_row.get(name, ""))
            except ValueError as exc:
                errors.append(ImportFieldError(field=name, message=str(exc)))

        persons: list[Person] = []
        pic_text = raw_row.get("pics", "")
        if pic_text and pic_text.upper() != "TBD":
            for pic_name in (n.strip() for n in re.split(r"[;,]", pic_text) if n.strip()):
                person = people.get(pic_name.lower())
                if person is None:
                    errors.append(
                        ImportFieldError(field="pics", message=f"Unknown PIC '{pic_name}'")
                    )
                elif not person.is_active:
                    errors.append(
                        ImportFieldError(field="pics", message=f"PIC '{pic_name}' is inactive")
                    )
                elif person not in persons:
                    persons.append(person)
        payload["pic_ids"] = [p.id for p in persons]

        create: ProjectCreate | None = None
        if not errors:
            try:
                create = ProjectCreate.model_validate(payload)
            except ValidationError as exc:
                for err in exc.errors():
                    loc = str(err["loc"][0]) if err["loc"] else None
                    msg = str(err["msg"]).removeprefix("Value error, ")
                    errors.append(ImportFieldError(field=loc, message=msg))

        name_key = raw_row.get("name", "").strip().lower()
        if name_key and not any(e.field == "name" for e in errors):
            if name_key in existing:
                errors.append(
                    ImportFieldError(
                        field="name", message="A project with this name already exists"
                    )
                )
            elif name_key in seen:
                errors.append(
                    ImportFieldError(
                        field="name", message=f"Duplicate of row {seen[name_key]} in this file"
                    )
                )
        if name_key:
            seen.setdefault(name_key, row_no)

        valid = not errors
        parsed.append((create, stage, persons) if valid and create and stage else None)
        report_rows.append(
            ImportRow(
                row=row_no,
                valid=valid,
                data={
                    "name": raw_row.get("name", ""),
                    "category": raw_row.get("category", ""),
                    "stage": stage_name,
                    **{f: raw_row.get(f, "") or None for f in DATE_FIELDS},
                    "pics": [p.name for p in persons],
                },
                errors=errors,
            )
        )
    return report_rows, parsed, ignored


def _report(
    rows: list[ImportRow], ignored: list[str], dry_run: bool, imported: int
) -> ImportReport:
    bad = [r for r in rows if not r.valid]
    return ImportReport(
        dry_run=dry_run,
        total_rows=len(rows),
        valid_rows=len(rows) - len(bad),
        error_rows=len(bad),
        imported=imported,
        ignored_columns=ignored,
        rows=rows,
        errors=[
            {"row": r.row, "field": e.field, "message": e.message} for r in bad for e in r.errors
        ],
    )


def import_csv(session: Session, raw: bytes, dry_run: bool) -> ImportReport:
    rows, parsed, ignored = _validate(session, raw)
    if dry_run:
        return _report(rows, ignored, True, 0)
    report = _report(rows, ignored, False, 0)
    if report.error_rows:
        raise AppError(
            422,
            "import_validation_failed",
            f"{report.error_rows} row(s) have errors. Nothing was imported.",
            report.model_dump(mode="json"),
        )
    if not rows:
        raise AppError(400, "invalid_csv", "The file has no data rows")
    for item in parsed:
        assert item is not None
        payload, stage, persons = item
        insert_project(session, payload, stage, persons)
    session.commit()
    return _report(rows, ignored, False, len(parsed))
