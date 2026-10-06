import csv
import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from freezegun import freeze_time
from openpyxl import load_workbook

from tests.conftest import Api

HEADER = (
    "Category,Project,Description,Stage,Status Notes,Plan Date,Adjusted Plan Date,"
    "Actual Date,Projected Start,PIC"
)
NOW = "2026-10-06 04:00:00"


def upload(client: TestClient, content: str | bytes, dry_run: bool | None = None) -> Any:
    data = content.encode() if isinstance(content, str) else content
    params = {} if dry_run is None else {"dry_run": str(dry_run).lower()}
    return client.post(
        "/api/v1/import/csv", params=params, files={"file": ("p.csv", data, "text/csv")}
    )


@pytest.fixture
def people(api: Api) -> None:
    api.person("Denan")
    api.person("Prada")
    api.person("Sleepy", is_active=False)


VALID = f"""{HEADER}
current,Alpha,First,In Progress,notes,2026-10-20,,,,Denan
pending,Beta,,Not Started,,,,,2026-11-01,Denan; Prada
maintenance,Gamma,,,,"Tue, Oct 06, 26",,,,TBD
"""


@freeze_time(NOW)
def test_dry_run_reports_without_writing(client: TestClient, api: Api, people: None) -> None:
    r = upload(client, VALID, dry_run=True)
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["dry_run"], body["total_rows"], body["valid_rows"], body["error_rows"]) == (
        True,
        3,
        3,
        0,
    )
    assert body["imported"] == 0
    assert [row["row"] for row in body["rows"]] == [2, 3, 4]
    assert body["rows"][1]["data"]["pics"] == ["Denan", "Prada"]
    assert api.names() == []  # nothing written


@freeze_time(NOW)
def test_real_import_creates_projects_with_activity(
    client: TestClient, api: Api, people: None
) -> None:
    r = upload(client, VALID)
    assert r.status_code == 201, r.text
    assert r.json()["imported"] == 3
    items = client.get("/api/v1/projects", params={"sort_by": "name"}).json()["items"]
    assert [i["name"] for i in items] == ["Alpha", "Beta", "Gamma"]
    alpha, beta, gamma = items
    assert alpha["stage"]["name"] == "In Progress" and alpha["plan_date"] == "2026-10-20"
    assert beta["projected_start_date"] == "2026-11-01"
    assert [p["name"] for p in beta["pics"]] == ["Denan", "Prada"]
    assert gamma["stage"]["name"] == "Not Started"  # default stage
    assert gamma["plan_date"] == "2026-10-06" and gamma["pics"] == []
    log = client.get(f"/api/v1/projects/{alpha['id']}/activity").json()
    assert [(e["field"], e["new_value"]) for e in log] == [("project", "created")]


@freeze_time(NOW)
def test_bad_rows_get_per_row_errors_and_block_real_import(
    client: TestClient, api: Api, people: None
) -> None:
    content = f"""{HEADER}
current,Good,,In Progress,,2026-10-20,,,,Denan
bogus,Bad Category,,In Progress,,,,,,
current,Bad Stage,,Nope,,,,,,
current,Bad Date,,In Progress,,31/12/2026,,,,
current,Future Actual,,In Progress,,,,2026-10-09,,
current,Bad Pic,,In Progress,,,,,,Nobody
current,Inactive Pic,,In Progress,,,,,,Sleepy
current,,,In Progress,,,,,,
current,Good,,In Progress,,,,,,
"""
    dry = upload(client, content, dry_run=True)
    assert dry.status_code == 200
    body = dry.json()
    assert body["total_rows"] == 9 and body["valid_rows"] == 1 and body["error_rows"] == 8
    by_row = {row["row"]: row for row in body["rows"]}
    assert by_row[2]["valid"] is True
    fields = {r: [e["field"] for e in by_row[r]["errors"]] for r in by_row if by_row[r]["errors"]}
    assert fields == {
        3: ["category"],
        4: ["stage"],
        5: ["plan_date"],
        6: ["actual_date"],
        7: ["pics"],
        8: ["pics"],
        9: ["name"],
        10: ["name"],
    }
    assert "Duplicate of row 2" in by_row[10]["errors"][0]["message"]
    assert {e["row"] for e in body["errors"]} == {3, 4, 5, 6, 7, 8, 9, 10}

    real = upload(client, content)
    assert real.status_code == 422
    err = real.json()["error"]
    assert err["code"] == "import_validation_failed"
    assert err["details"]["error_rows"] == 8
    assert api.names() == []  # all-or-nothing


def test_existing_name_is_an_error(client: TestClient, api: Api, people: None) -> None:
    api.project("Alpha")
    r = upload(client, f"{HEADER}\ncurrent,alpha,,,,,,,,\n", dry_run=True)
    row = r.json()["rows"][0]
    assert not row["valid"] and "already exists" in row["errors"][0]["message"]


def test_file_level_errors_are_400(client: TestClient) -> None:
    assert upload(client, "").status_code == 400
    assert upload(client, b"\xff\xfe\x00bad").status_code == 400
    r = upload(client, "Foo,Bar\n1,2\n")
    assert r.status_code == 400 and "Project" in r.json()["error"]["message"]
    r = upload(client, "Category,Project,Project\ncurrent,A,B\n")
    assert r.status_code == 400
    r = upload(client, f"{HEADER}\n", dry_run=True)
    assert r.status_code == 200 and r.json()["total_rows"] == 0
    r = upload(client, f"{HEADER}\n", dry_run=False)
    assert r.status_code == 400  # real import needs at least one row
    assert client.post("/api/v1/import/csv").status_code == 422  # no file


def test_size_and_row_limits(client: TestClient) -> None:
    assert upload(client, b"a" * (2 * 1024 * 1024 + 10)).status_code == 413
    rows = "\n".join(f"current,P{i},,,,,,,," for i in range(2001))
    assert upload(client, f"{HEADER}\n{rows}\n").status_code == 400


def test_semicolon_delimiter_blank_lines_and_ignored_columns(client: TestClient, api: Api) -> None:
    content = "Category;Project;Extra Column;Stage\ncurrent;A;x;Testing\n;;;\ncurrent;B;y;Done\n"
    r = upload(client, content, dry_run=True)
    body = r.json()
    assert body["total_rows"] == 2 and body["valid_rows"] == 2
    assert body["ignored_columns"] == ["Extra Column"]
    assert [row["row"] for row in body["rows"]] == [2, 4]  # blank line skipped, numbering kept


def test_quoted_multiline_cells_and_bom(client: TestClient, api: Api) -> None:
    content = '﻿Category,Project,Description\ncurrent,"A, B","line1\nline2"\n'
    r = upload(client, content)
    assert r.status_code == 201
    item = client.get("/api/v1/projects").json()["items"][0]
    assert item["name"] == "A, B" and item["description"] == "line1\nline2"


@freeze_time(NOW)
def test_export_csv_uses_filters_and_display_dates(client: TestClient, api: Api) -> None:
    ann = api.person("Ann")
    api.project("Alpha", plan_date="2026-10-01", pic_ids=[ann["id"]], description="=SUM(A1)")
    api.project("Beta", category="pending")
    r = client.get("/api/v1/export", params={"format": "csv", "category": "current"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "projects-2026-10-06.csv" in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["Project"] for x in rows] == ["Alpha"]
    row = rows[0]
    assert row["Plan Date"] == "Thu, Oct 01, 26"
    assert row["Adjusted Plan Date"] == "TBD"
    assert row["PIC"] == "Ann" and row["Deadline Status"] == "Overdue"
    assert row["Description"] == "'=SUM(A1)"  # formula injection neutralised


@freeze_time(NOW)
def test_export_xlsx(client: TestClient, api: Api) -> None:
    api.project("Alpha", plan_date="2026-10-01")
    r = client.get("/api/v1/export", params={"format": "xlsx"})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    headers = [c.value for c in ws[1]]
    assert headers[:2] == ["Category", "Project"]
    row = {h: c.value for h, c in zip(headers, ws[2], strict=True)}
    assert row["Project"] == "Alpha"
    assert str(row["Plan Date"])[:10] == "2026-10-01"
    assert ws.cell(row=2, column=headers.index("Plan Date") + 1).number_format == "ddd, mmm dd, yy"
    assert row["Adjusted Plan Date"] == "TBD"


def test_export_invalid_format_and_filter_parity(client: TestClient, api: Api) -> None:
    assert client.get("/api/v1/export", params={"format": "pdf"}).status_code == 400
    api.project("Alpha")
    api.project("Beta")
    r = client.get("/api/v1/export", params={"search": "alp", "sort_by": "name"})
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["Project"] for x in rows] == ["Alpha"]


@freeze_time(NOW)
def test_export_then_import_round_trip(client: TestClient, api: Api) -> None:
    ann = api.person("Ann")
    api.project(
        "Alpha",
        plan_date="2026-10-01",
        adjusted_plan_date="2026-10-09",
        pic_ids=[ann["id"]],
        description="=x",
    )
    api.project(
        "Beta", category="pending", projected_start_date="2026-11-02", status_notes="hello, world"
    )
    exported = client.get("/api/v1/export", params={"sort_by": "name"}).content
    for p in client.get("/api/v1/projects").json()["items"]:
        client.delete(f"/api/v1/projects/{p['id']}")
    r = upload(client, exported)
    assert r.status_code == 201, r.text
    items = {i["name"]: i for i in client.get("/api/v1/projects").json()["items"]}
    assert items["Alpha"]["adjusted_plan_date"] == "2026-10-09"
    assert items["Alpha"]["description"] == "=x"
    assert items["Alpha"]["pics"][0]["name"] == "Ann"
    assert items["Beta"]["projected_start_date"] == "2026-11-02"
    assert items["Beta"]["status_notes"] == "hello, world"
